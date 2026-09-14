from pathlib import Path
import re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
p=Path(__file__).resolve().parents[2]
f=p/"04_申报与答辩/申报书_教师反馈主线.md"
s=f.read_text(encoding="utf-8-sig")
doc=Document()
sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.3);sec.bottom_margin=Cm(2.2);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.5)
for name in ["Normal","Title","Heading 1","Heading 2","Heading 3","Caption"]:
 st=doc.styles[name];st.font.name="Times New Roman";st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"宋体" if name in ["Normal","Caption"] else "黑体")
 st.font.size=Pt(12 if name=="Normal" else 14)
 st.paragraph_format.space_after=Pt(5)
 st.paragraph_format.line_spacing=1.35
doc.styles["Title"].font.size=Pt(21)
doc.styles["Heading 1"].font.size=Pt(15)
doc.styles["Heading 1"].paragraph_format.space_before=Pt(14)
doc.styles["Heading 2"].font.size=Pt(12)
doc.styles["Heading 2"].paragraph_format.space_before=Pt(9)
doc.styles["Caption"].font.size=Pt(10)
for st in doc.styles:
 for el in list(st.element.iter(qn("w:pBdr"))): el.getparent().remove(el)
doc.core_properties.title="大连理工大学大学生创新训练项目申报书"
doc.core_properties.subject="面向录播网课的弹幕驱动多模态教学反馈解析与教师复盘系统"
doc.core_properties.author=""
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run();fld=OxmlElement("w:fldSimple");fld.set(qn("w:instr"),"PAGE");r._r.addnext(fld)
def fmt(par,body=True):
 par.paragraph_format.widow_control=True
 if body:par.paragraph_format.first_line_indent=Pt(24)
def addtable(lines):
 rows=[[v.strip() for v in l.strip().strip("|").split("|")] for l in lines if not re.match(r"^\|[\s:|\-]+\|$",l)]
 t=doc.add_table(rows=0,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=[5.1,2.5,8.4]
 for j,col in enumerate(t.columns): col.width=Cm(widths[j])
 for ir,row in enumerate(rows):
  cells=t.add_row().cells
  for j,val in enumerate(row):
   cells[j].width=Cm(widths[j] if len(row)==3 else 16/len(row));cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=cells[j].paragraphs[0];cp.paragraph_format.line_spacing=1.15;cp.paragraph_format.space_before=Pt(5);cp.paragraph_format.space_after=Pt(5)
   cp.add_run(val).font.size=Pt(10.5)
   if j==1:cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
   pr=cells[j]._tc.get_or_add_tcPr()
   borders=OxmlElement("w:tcBorders")
   for edge in ["top","left","bottom","right"]:
    el=OxmlElement("w:"+edge);el.set(qn("w:val"),"single");el.set(qn("w:sz"),"4");el.set(qn("w:color"),"D9D9D9");borders.append(el)
   pr.append(borders)
   if ir==0:
    shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"EEEEEE");pr.append(shade)
    for run in cp.runs:run.bold=True
  trpr=t.rows[-1]._tr.get_or_add_trPr();cant=OxmlElement("w:cantSplit");trpr.append(cant)
  if ir<=1:
   for cell in cells:
    for cp in cell.paragraphs:cp.paragraph_format.keep_with_next=True
  if ir==0:
   trpr.append(OxmlElement("w:tblHeader"))
   for cell in cells:
    for cp in cell.paragraphs:cp.paragraph_format.keep_with_next=True
 doc.add_paragraph().paragraph_format.space_after=Pt(2)
lines=s.splitlines();i=0;reference=False
while i<len(lines):
 l=lines[i].strip();i+=1
 if not l:continue
 if l.startswith("|"):
  table=[l]
  while i<len(lines) and lines[i].strip().startswith("|"):table.append(lines[i]);i+=1
  addtable(table);continue
 if l.startswith("!["):
  target=re.search(r"\]\((.*?)\)",l).group(1)
  par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER;par.paragraph_format.keep_with_next=True
  par.add_run().add_picture(str(f.parent/target),width=Cm(14.0 if "研究技术路线" in target else 15.7))
  continue
 if l.startswith("# "):
  par=doc.add_paragraph(l[2:],style="Title");par.alignment=WD_ALIGN_PARAGRAPH.CENTER
  par=doc.add_paragraph("2027年");par.alignment=WD_ALIGN_PARAGRAPH.CENTER
  continue
 if l.startswith("## "):
  title=l[3:].replace("、"," ")
  par=doc.add_paragraph(title,style="Heading 1")
  reference="主要参考文献" in title
  if title.startswith(("三 ",)):par.paragraph_format.page_break_before=True
  continue
 if l.startswith("### "):
  doc.add_paragraph(l[4:],style="Heading 2");continue
 if re.match(r"^图[12] ",l):
  par=doc.add_paragraph(l,style="Caption");par.alignment=WD_ALIGN_PARAGRAPH.CENTER;par.paragraph_format.keep_with_next=False;continue
 par=doc.add_paragraph(l)
 basic=(l.startswith(("项目名称：","项目类型：","申报单位：","申报批次：","学科方向：","项目来源：","负责人学号","项目负责人：","团队成员：","指导教师：","执行时间：","指导教师签名","负责人签名","___")))
 fmt(par,not basic and not reference)
 if reference:
  par.paragraph_format.line_spacing=1.1
  par.paragraph_format.space_after=Pt(5)
  for run in par.runs:run.font.size=Pt(9.5)
output=f.with_suffix(".docx");doc.save(output)
print(output)

